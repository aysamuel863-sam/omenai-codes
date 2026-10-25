from docx import Document

document = Document()

document.add_heading("Photosynthesis", level = 1)
document.add_paragraph("Photosynthesis is the process by " \
"which green plants take food using sunlight, water and carbon dioxide")
document.add_paragraph("It occurs in plants. It does not occur in animals")

document.save("notes.docx")