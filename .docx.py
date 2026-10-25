import os 
import json
from openai import OpenAI
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import tkinter as tk
from tkinter import scrolledtext

client = OpenAI(
    base_url = "https://integrate.api.nvidia.com/v1",
    api_key = os.getenv("NVIDIA_API_KEY2")
)

root = tk.Tk()
root.title("Write2Doc")
root.geometry("600x450")

label = tk.Label(root, text = "Paste your document below:")
label.pack(pady = 100)

input_box = scrolledtext.ScrolledText(
    root,
    wrap = tk.WORD,
    width = 65,
    height = 18
)
input_box.pack(padx = 10, pady = 10)

button = tk.Button(
    root,
    text = "Convert to Word",
    command = root.quit
)
button.pack(pady = 10)

root.mainloop()
text = input_box.get("1.0", "end-1c")

root.destroy()
prompt = f"""
You are Write2Doc, an AI document formattinng assistant.
Read the user's text and organize it into a clear,
well-structured and professional document.
Analyze the content yourself and decide the most
appropriate structure and formatting.
Return ONLY valid JSON.
Use this structure:
{{
    "title": {{
        "text": "Document title",
        "bold": true,
        "italic": false,
        "size": 20,
        "alignment": "center"               
    }}
    "sections": [
    {{
        "heading": {{
            "text": "Section heading",
            "level": 2,
            "bold": true,
            "italic": false,
            "size": 16
        }}
        "type": "paragraph",
        "content": "Paragraph text"
    }}
    ]
}}
For every section, independently choose the type
that best fits the information.
Allowed section types:
- paragraph
- bullet_list
- numbered_list
- table
For paragraph:
"content" must be a string.
For bullet_list:
"content" must be an array of strings.
For numbered_list:
"content" must be an array of strings.
For table:
"content" must be a two-dimensional array.
Choose appropriate formatting based on the content.
For titles and headings, choose appropriate:
- bold
- italic
- size
- alignment
Do not invent information.
Preserve the meaning of the original text.
Do not use unnecessary formatting.
User's text:
{text}
"""

message = [{
        "role": "user",
        "content": prompt
}]
response = client.chat.completions.create(
        model = "nvidia/nemotron-3-super-120b-a12b",
        messages = [{
        "role": "user",
        "content": prompt
}],
    #response_format = {"type": "json_object"},

        max_tokens = 1500,
        extra_body = {"chat_template_kwargs" : {"enable_thinking": "False"}}
        #temperature = 0.5,
        #top_p = 1,
        #stream = False
)
#print(response.choices[0].message.content)

ai_output = response.choices[0].message.content

print("\nAI:")
print(ai_output)

data = json.loads(ai_output)
document = Document()
title = document.add_paragraph()
title_run = title.add_run(data["title"]["text"])
title_run.bold = data["title"]["bold"]
title_run.italic = data["title"]["italic"]
title_run.font.size = Pt(
    data["title"]["size"]
)
if data["title"]["alignment"] == "center":
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
elif data["title"]["alignment"] == "right":
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
else:
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
for section in data["sections"]:
    heading = document.add_paragraph()
    heading_run = heading.add_run(
    section["heading"]["text"]
    )
    heading_run.bold = section["heading"]["bold"]
    heading_run.italic = section["heading"]["italic"]
    heading_run.font.size = Pt(
    section["heading"]["size"]
    )
    section_type = section["type"]
    content = section["content"]
    if section_type == "paragraph":
        document.add_paragraph(content)
    elif section_type == "bullet_list":
        for item in content:
            document.add_paragraph(
            item,
            style="List Bullet"
            )
    elif section_type == "numbered_list":
        for item in content:
            document.add_paragraph(
            item,
            style="List Number"
            )
    elif section_type == "table":
        rows = len(content)
        columns = len(content[0])
        table = document.add_table(
        rows=rows,
        cols=columns
        )
        for row_index, row in enumerate(content):
            for column_index, value in enumerate(row):
                table.cell(
                row_index,
                column_index
                ).text = str(value)
document.save("write2doc_test.docx")
print("\nDocument created successfully!")